"""
HackersJack Cybersecurity Education
Statistical Reproducibility Pipeline

Study:
Evaluating HackersJack Cybersecurity Education:
A Pre-Post Study of Grade 5 and 6 Students in Australia

Purpose:
Reproduce the statistical analyses using the original
12 Grade 5 and Grade 6 Excel assessment workbooks.

Input directory:
    data/raw_assessment_data/

Output directory:
    outputs/

Outputs:
    statistical_results.csv
    verification.csv
    data_quality_counts.csv
    source_checksums.json

No individual student-level records are exported.
"""

import argparse
import csv
import hashlib
import json
import re
import sys

from collections import defaultdict
from pathlib import Path
from statistics import mean, stdev

from openpyxl import load_workbook
from scipy import stats


# ============================================================
# 1. STUDY CONFIGURATION
# ============================================================

OUTCOMES = ("MCQ", "TF", "Combined")

EXPECTED = {
    "Combined": {
        "n": 216,
        "pre_mean": 21.61,
        "post_mean": 28.67,
        "gain": 7.06,
        "t": 40.17,
        "dz": 2.73,
    },
    "MCQ": {
        "n": 225,
        "pre_mean": 14.61,
        "post_mean": 19.09,
        "gain": 4.48,
        "t": 31.49,
        "dz": 2.10,
    },
    "TF": {
        "n": 231,
        "pre_mean": 7.04,
        "post_mean": 9.58,
        "gain": 2.53,
        "t": 28.50,
        "dz": 1.88,
    },
    "Grade 5": {
        "n": 107,
        "pre_mean": 22.44,
        "post_mean": 28.81,
        "gain": 6.37,
        "t": 28.46,
        "dz": 2.75,
    },
    "Grade 6": {
        "n": 109,
        "pre_mean": 20.79,
        "post_mean": 28.52,
        "gain": 7.73,
        "t": 30.33,
        "dz": 2.91,
    },
}


# ============================================================
# 2. READ ASSESSMENT WORKSHEETS
# ============================================================

def parse_sheet(sheet, kind, grade, class_id, file_name):
    """
    Extract pre- and post-assessment records.

    MCQ: 20 binary questions.
    TF:  10 binary questions.

    Correct answer = 1
    Incorrect answer = 0

    Missing or invalid answers invalidate the assessment
    occasion for matched-pair analysis.
    """

    rows = list(
        sheet.iter_rows(
            min_row=1,
            max_row=150,
            min_col=1,
            max_col=22,
            values_only=True,
        )
    )

    sections = []

    for i, row in enumerate(rows):

        cell = str(row[0] or "").strip().lower()

        if re.fullmatch(r"(pre|post)\s*assessment", cell):

            occasion = (
                "pre" if cell.startswith("pre") else "post"
            )

            sections.append((i, occasion))

    if [s[1] for s in sections] != ["pre", "post"]:

        raise ValueError(
            f"{file_name} / {sheet.title}: "
            f"expected Pre and Post sections; "
            f"found {sections}"
        )

    result = []

    nitems = 20 if kind == "MCQ" else 10

    for k, (start, occasion) in enumerate(sections):

        stop = (
            sections[k + 1][0]
            if k + 1 < len(sections)
            else len(rows)
        )

        block = rows[start + 1:stop]

        header = next(
            (
                j
                for j, row in enumerate(block)
                if str(row[1] or "").strip().upper() == "Q1"
            ),
            None,
        )

        if header is None:

            raise ValueError(
                f"{file_name}/{sheet.title}/{occasion}: "
                "missing Q1 header"
            )

        seen = set()

        for row in block[header + 1:]:

            code = str(row[0] or "").strip()

            if not code:
                continue

            if code.lower() in (
                "total",
                "average",
                "mean",
                "sum",
            ):
                continue

            if not re.search(r"\d", code):
                continue

            if code in seen:

                raise ValueError(
                    f"Duplicate student code: "
                    f"{file_name}/{sheet.title}/"
                    f"{occasion}/{code}"
                )

            seen.add(code)

            vals = row[1:nitems + 1]

            valid = all(
                isinstance(v, (int, float))
                and not isinstance(v, bool)
                and v in (0, 1)
                for v in vals
            )

            score = (
                int(sum(vals))
                if valid
                else None
            )

            result.append({
                "grade": grade,
                "class": class_id,
                "code": code,
                "kind": kind,
                "occasion": occasion,
                "score": score,
                "valid": valid,
                "items": vals if valid else None,
                "file": file_name,
            })

    return result


# ============================================================
# 3. LOAD ORIGINAL EXCEL DATASETS
# ============================================================

def read_records(folder):
    """
    Read all 12 original HackersJack Excel workbooks.

    Expected:
        6 Grade 5 workbooks
        6 Grade 6 workbooks

    Each workbook contains:
        1 MCQ worksheet
        1 True/False worksheet
    """

    folder = Path(folder)

    if not folder.exists():

        raise FileNotFoundError(
            f"Dataset directory not found: "
            f"{folder.resolve()}"
        )

    files = sorted(
        path
        for path in folder.glob("*.xlsx")
        if not path.name.startswith("~$")
    )

    if len(files) != 12:

        raise ValueError(
            f"Expected 12 Excel workbooks, "
            f"found {len(files)} in {folder.resolve()}"
        )

    records = []
    hashes = []

    grade_counts = defaultdict(int)

    for path in files:

        print(f"Reading: {path.name}")

        grade_match = re.search(
            r"Grade\s*([56])",
            path.name,
            re.IGNORECASE,
        )

        if not grade_match:

            raise ValueError(
                f"Cannot identify grade: {path.name}"
            )

        grade = int(grade_match.group(1))

        grade_counts[grade] += 1

        class_id = path.name.split(
            "Pre and Post"
        )[0].strip(" -")

        sha256 = hashlib.sha256(
            path.read_bytes()
        ).hexdigest()

        hashes.append({
            "filename": path.name,
            "sha256": sha256,
        })

        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )

        try:

            sheets = workbook.worksheets

            if len(sheets) != 2:

                raise ValueError(
                    f"{path.name}: expected 2 worksheets, "
                    f"found {len(sheets)}"
                )

            for sheet in sheets:

                name = sheet.title.upper()

                if "MCQ" in name:

                    kind = "MCQ"

                elif (
                    "T-F" in name
                    or "T_F" in name
                ):

                    kind = "TF"

                else:

                    raise ValueError(
                        f"Unknown assessment type: "
                        f"{path.name}/{sheet.title}"
                    )

                parsed = parse_sheet(
                    sheet,
                    kind,
                    grade,
                    class_id,
                    path.name,
                )

                records.extend(parsed)

        finally:

            workbook.close()

    if grade_counts[5] != 6 or grade_counts[6] != 6:

        raise ValueError(
            "Expected six Grade 5 and six Grade 6 "
            f"workbooks; found {dict(grade_counts)}"
        )

    print(
        f"\nLoaded {len(files)} workbooks "
        f"and {len(records)} assessment records."
    )

    return records, hashes


# ============================================================
# 4. CONSTRUCT MATCHED PRE/POST PAIRS
# ============================================================

def construct_pairs(records):
    """
    Match students by:
        grade
        class
        student code
        assessment component

    Combined scores require valid matched observations
    for both MCQ and True/False components.
    """

    bykey = defaultdict(dict)

    for record in records:

        key = (
            record["grade"],
            record["class"],
            record["code"],
            record["kind"],
        )

        occasion = record["occasion"]

        if occasion in bykey[key]:

            raise ValueError(
                f"Duplicate assessment key: "
                f"{key}/{occasion}"
            )

        bykey[key][occasion] = record

    component = {}

    for key, observations in bykey.items():

        if all(
            occasion in observations
            and observations[occasion]["valid"]
            for occasion in ("pre", "post")
        ):

            component[key] = {
                "grade": key[0],
                "class": key[1],
                "code": key[2],
                "kind": key[3],
                "pre": observations["pre"]["score"],
                "post": observations["post"]["score"],
            }

    combined = []

    for (grade, class_id, code, kind), record in component.items():

        if kind != "MCQ":
            continue

        tf = component.get(
            (grade, class_id, code, "TF")
        )

        if tf is not None:

            combined.append({
                "grade": grade,
                "class": class_id,
                "code": code,
                "kind": "Combined",
                "pre": record["pre"] + tf["pre"],
                "post": record["post"] + tf["post"],
            })

    return list(component.values()) + combined


# ============================================================
# 5. STATISTICAL ANALYSIS
# ============================================================

def summarize(pairs):
    """
    Calculate paired-samples statistics.

    Includes:
        N
        Pre/post means
        Pre/post standard deviations
        Mean gain
        Paired t-test
        Degrees of freedom
        p-value
        Cohen's dz
        95% confidence interval
    """

    if len(pairs) < 2:
        return None

    pre = [
        record["pre"]
        for record in pairs
    ]

    post = [
        record["post"]
        for record in pairs
    ]

    differences = [
        post_score - pre_score
        for pre_score, post_score in zip(pre, post)
    ]

    n = len(differences)

    gain_mean = mean(differences)

    gain_sd = stdev(differences)

    if gain_sd == 0:

        t_statistic = float("nan")

        dz = float("nan")

        ci = (
            float("nan"),
            float("nan"),
        )

        p_value = float("nan")

    else:

        standard_error = gain_sd / (n ** 0.5)

        t_statistic = gain_mean / standard_error

        dz = gain_mean / gain_sd

        ci = stats.t.interval(
            confidence=0.95,
            df=n - 1,
            loc=gain_mean,
            scale=standard_error,
        )

        p_value = float(
            2 * stats.t.sf(
                abs(t_statistic),
                n - 1,
            )
        )

    return {
        "n": n,
        "pre_mean": mean(pre),
        "pre_sd": stdev(pre),
        "post_mean": mean(post),
        "post_sd": stdev(post),
        "gain": gain_mean,
        "gain_sd": gain_sd,
        "t": t_statistic,
        "df": n - 1,
        "p": p_value,
        "dz": dz,
        "ci_low": ci[0],
        "ci_high": ci[1],
    }


# ============================================================
# 6. REPRODUCIBILITY VERIFICATION
# ============================================================

def verify_results(groups):
    """
    Compare recalculated statistics against
    manuscript values at two-decimal precision.
    """

    validation = []

    for outcome, expected in EXPECTED.items():

        observed = groups[outcome]

        if observed is None:

            raise ValueError(
                f"No matched observations: {outcome}"
            )

        for metric, reported in expected.items():

            recalculated = observed[metric]

            if metric == "n":

                match = recalculated == reported

            else:

                match = (
                    abs(recalculated - reported)
                    < 0.00500001
                )

            validation.append({
                "outcome": outcome,
                "metric": metric,
                "reported": reported,
                "recomputed": recalculated,
                "rounding_match": match,
            })

    return validation


# ============================================================
# 7. RUN COMPLETE ANALYSIS
# ============================================================

def run(input_dir, output_dir):

    input_dir = Path(input_dir)

    output_dir = Path(output_dir)

    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    print("\n====================================")
    print("HackersJack Statistical Verification")
    print("====================================\n")

    records, hashes = read_records(input_dir)

    pairs = construct_pairs(records)

    groups = {}

    for kind in OUTCOMES:

        groups[kind] = summarize([
            record
            for record in pairs
            if record["kind"] == kind
        ])

    for grade in (5, 6):

        groups[f"Grade {grade}"] = summarize([
            record
            for record in pairs
            if record["kind"] == "Combined"
            and record["grade"] == grade
        ])

    validation = verify_results(groups)

    # --------------------------------------------------------
    # Export aggregate statistical results
    # --------------------------------------------------------

    statistical_file = (
        output_dir / "statistical_results.csv"
    )

    with statistical_file.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as file:

        fieldnames = [
            "outcome",
            *next(iter(groups.values())).keys(),
        ]

        writer = csv.DictWriter(
            file,
            fieldnames=fieldnames,
        )

        writer.writeheader()

        for outcome, result in groups.items():

            writer.writerow({
                "outcome": outcome,
                **result,
            })

    # --------------------------------------------------------
    # Export verification report
    # --------------------------------------------------------

    verification_file = (
        output_dir / "verification.csv"
    )

    with verification_file.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=[
                "outcome",
                "metric",
                "reported",
                "recomputed",
                "rounding_match",
            ],
        )

        writer.writeheader()

        writer.writerows(validation)

    # --------------------------------------------------------
    # Export source checksums
    # --------------------------------------------------------

    checksum_file = (
        output_dir / "source_checksums.json"
    )

    with checksum_file.open(
        "w",
        encoding="utf-8",
    ) as file:

        json.dump(
            hashes,
            file,
            indent=2,
        )

    # --------------------------------------------------------
    # Export data-quality counts
    # --------------------------------------------------------

    counts = defaultdict(int)

    for record in records:

        status = (
            "valid"
            if record["valid"]
            else "unavailable"
        )

        key = (
            record["kind"],
            record["occasion"],
            status,
        )

        counts[key] += 1

    quality_file = (
        output_dir / "data_quality_counts.csv"
    )

    with quality_file.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as file:

        writer = csv.writer(file)

        writer.writerow([
            "component",
            "occasion",
            "status",
            "records",
        ])

        for (
            kind,
            occasion,
            status,
        ), count in sorted(counts.items()):

            writer.writerow([
                kind,
                occasion,
                status,
                count,
            ])

    # --------------------------------------------------------
    # Display verification summary
    # --------------------------------------------------------

    all_match = all(
        item["rounding_match"]
        for item in validation
    )

    print("\nSTATISTICAL RESULTS\n")

    for outcome, result in groups.items():

        print(
            f"{outcome}: "
            f"N={result['n']}, "
            f"Pre={result['pre_mean']:.2f}, "
            f"Post={result['post_mean']:.2f}, "
            f"Gain={result['gain']:.2f}, "
            f"t={result['t']:.2f}, "
            f"dz={result['dz']:.2f}"
        )

    print("\nVERIFICATION RESULTS\n")

    for item in validation:

        status = (
            "PASS"
            if item["rounding_match"]
            else "FAIL"
        )

        print(
            f"{status}: "
            f"{item['outcome']} / "
            f"{item['metric']}"
        )

    print("\n====================================")

    if all_match:

        print("ALL STATISTICAL VERIFICATION CHECKS PASSED")

    else:

        print("SOME STATISTICAL VERIFICATION CHECKS FAILED")

    print("====================================\n")

    print(
        f"Reports saved to: "
        f"{output_dir.resolve()}"
    )

    return all_match


# ============================================================
# 8. COMMAND-LINE EXECUTION
# ============================================================

if __name__ == "__main__":

    parser = argparse.ArgumentParser(
        description=(
            "Reproduce HackersJack "
            "pre/post statistical analyses"
        )
    )

    parser.add_argument(
        "--input",
        type=Path,
        default=Path("data/raw_assessment_data"),
        help="Directory containing 12 Excel workbooks",
    )

    parser.add_argument(
        "--output",
        type=Path,
        default=Path("outputs"),
        help="Directory for statistical reports",
    )

    args = parser.parse_args()

    success = run(
        input_dir=args.input,
        output_dir=args.output,
    )

    sys.exit(
        0 if success else 2
    )
